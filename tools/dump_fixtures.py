"""Regenerates test/fixtures/*.json from the real calibration workbooks.

The harness mocks Excel with plain JSON: for each sheet, the values grid, the used-range offset,
and (for the PM Fee workbook) the formulas grid, where a constant cell carries its literal value,
exactly as ExcelScript getFormulas() returns it.

Run:  pip install openpyxl  &&  python tools/dump_fixtures.py
Add a new source: put the file in calibration/, add a line to FIXTURES, re-run, then load it in a test.
"""
import datetime, json, os, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CAL = os.path.join(ROOT, "calibration", "610West_07-2026")
OUT = os.path.join(ROOT, "test", "fixtures")

# fixture name -> (source file, include formulas?)
FIXTURES = {
    "tb":    ("Trial_Balance_mn610wes_accrual.xlsx", False),
    "bsss":  ("07_2026_BS_Support_Schedule__-_610_West.xlsm", False),
    "t12":   ("12_Month_Statement_mn610wes_accrual__3_.xlsx", False),
    "bc":    ("Budget_Comparison_mn610wes_accrual.xlsx", False),
    "bs":    ("Balance_Sheet_PeriodChange_mn610wes_accrual__13_.xlsx", False),
    "gl":    ("General_Ledger_mn610wes_accrual.xlsx", False),
    "ar":    ("AgedReceivables07_24_2026__1_.xlsx", False),
    "ap":    ("PayablesAgingReport07_24_2026__1_.xlsx", False),
    "gpr":   ("GrossPotentialRent07_24_2026__2_.xlsx", False),
    "sd":    ("SecurityDepositActivity07_24_2026__1_.xlsx", False),
    "rr":    ("RentRollwithLeaseCharges07_24_2026__2_.xlsx", False),
    "pmfee": ("07_2026_Yardi_PM_Fee_-_Fee_Pools_w_Custom_Report-610_West.xlsm", True),
}

def conv(v):
    if v is None: return ""
    if isinstance(v, datetime.datetime): return (v - datetime.datetime(1899, 12, 30)).total_seconds() / 86400
    if isinstance(v, datetime.date): return (datetime.datetime(v.year, v.month, v.day) - datetime.datetime(1899, 12, 30)).days
    return v if isinstance(v, (int, float, str, bool)) else str(v)

def dump(path, with_formulas):
    wv = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path) if with_formulas else None
    out = {}
    for ws in wv.worksheets:
        sheet = {"r0": ws.min_row - 1, "c0": ws.min_column - 1,
                 "vals": [[conv(v) for v in r] for r in ws.iter_rows(values_only=True)]}
        if wf:
            f = wf[ws.title]
            sheet["forms"] = [[(c.value if isinstance(c.value, str) and c.value.startswith("=") else conv(vc.value))
                               for c, vc in zip(rf, rv)] for rf, rv in zip(f.iter_rows(), ws.iter_rows())]
        out[ws.title] = sheet
    return out

if __name__ == "__main__":
    if "--template" in sys.argv:
        t = os.path.join(ROOT, "template", "Internal_Review_Report_Template_v2_0.xlsx")
        with open(os.path.join(OUT, "template.json"), "w") as fh: json.dump(dump(t, False), fh)
        print("wrote test/fixtures/template.json"); sys.exit(0)
    only = set(sys.argv[1:])
    for name, (fname, forms) in FIXTURES.items():
        if only and name not in only: continue
        src = os.path.join(CAL, fname)
        with open(os.path.join(OUT, name + ".json"), "w") as fh: json.dump(dump(src, forms), fh)
        print(f"wrote test/fixtures/{name}.json  <- {fname}")
    print("pdf_bankrec.json and pdf_mortgage.json are hand-transcribed from the PDFs (what the AI prompts should return); edit them directly.")
    print("template.json is the report template; regenerate it with:  python tools/dump_fixtures.py --template")
