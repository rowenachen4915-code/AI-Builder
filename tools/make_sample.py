"""Rebuilds template/SAMPLE_mn610wes_07.2026_Internal_Review_R1.xlsx from a calibration run.

  npm run calibrate            (writes build/filled.json)
  python tools/make_sample.py

It copies the report template, writes every value the review engine wrote in the mock workbook,
skips cells that hold formulas in the template (the Executive Summary counts), and adds the
All Results table. Formulas recalculate when the file is opened in Excel.
"""
import json, os, shutil
import openpyxl
from openpyxl.utils import range_boundaries

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "template", "Internal_Review_Report_Template_v2_0.xlsx")
OUT = os.path.join(ROOT, "template", "SAMPLE_mn610wes_07.2026_Internal_Review_R1.xlsx")

filled = json.load(open(os.path.join(ROOT, "build", "filled.json")))
orig = json.load(open(os.path.join(ROOT, "test", "fixtures", "template.json")))
shutil.copy(TEMPLATE, OUT)
wb = openpyxl.load_workbook(OUT)
for name, sheet in filled.items():
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    base = orig.get(name, {"vals": []})["vals"]
    for i, row in enumerate(sheet["vals"]):
        for j, v in enumerate(row):
            was = base[i][j] if i < len(base) and j < len(base[i]) else ""
            if v != "" and v != was:
                c = ws.cell(i + 1, j + 1)
                if isinstance(c.value, str) and c.value.startswith("="): continue
                c.value = v
# Grow any template table the engine outran (the script does this with table.resize in Excel).
# Count only the unbroken run of rows under the header: on Source Validation the grid table is
# followed by a blank row and a second table, and growing into it makes Excel repair the file.
for ws in wb.worksheets:
    for t in ws.tables.values():
        first, last = t.ref.split(":")
        col = "".join(ch for ch in last if ch.isalpha())
        top = int("".join(ch for ch in first if ch.isdigit()))
        used = top
        while ws.cell(used + 1, 1).value not in (None, ""): used += 1
        if used > int("".join(ch for ch in last if ch.isdigit())):
            t.ref = f"{first}:{col}{used}"
            if t.autoFilter is not None: t.autoFilter.ref = t.ref
# Excel rejects overlapping tables, so refuse to write a file that has them.
for ws in wb.worksheets:
    boxes = [(t.name, range_boundaries(t.ref)) for t in ws.tables.values()]
    for i, (na, a) in enumerate(boxes):
        for nb, b in boxes[i + 1:]:
            if not (a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1]):
                raise SystemExit(f"{ws.title}: table {na} overlaps {nb}")
wb.save(OUT)
print("wrote", OUT)
