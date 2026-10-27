"""Builds template/Internal_Review_Report_Template_v2_0.xlsx from v1.2.

v2.0 changes (requested by the owner, 2026-09-26):
  - Findings are split by review area instead of by status: BSSS, Bank Rec, T12, Budget Comparison,
    plus General Ledger, PM Fee and Other Controls so every finding has a home. Each tab lists
    Exceptions first, then Warnings, Monitor and Pass.
  - Column order puts what a manager reads first on screen: Status, Account / Item, Finding,
    Required Action. Test Performed moves right. Rule ID, Result ID, Property Code and Review
    Period are hidden (the same on every row, or only needed by the tooling), not deleted.
  - The review no longer decides readiness. Executive Summary D8 and Review History column D are
    a Manager Decision dropdown (Ready for Release / Ready With Comments / Not Ready), left blank.
  - Executive Summary gains a Review Areas table, a BSSS section (Ties / Does Not Tie / Missing
    Information / Supported Elsewhere, plus the accounts needing attention) and a Bank
    Reconciliations section (Clean / Monitor / Warning / Exception per rec).

Kept from v1.2: Calibri, navy section bands, byline, Source Validation grid, Changes Since Last
Review, Review History, and the hidden Setup and Rule Library sheets.

THE LAYOUT HERE IS A CONTRACT WITH scripts/04_RunInternalReview.ts. Row and column positions
marked "engine writes" must match the constants in that script's report writer.

Run:  python tools/build_template.py     then   python tools/dump_fixtures.py --template
"""
import os
import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "template", "Internal_Review_Report_Template_v1_2.xlsx")
OUT = os.path.join(ROOT, "template", "Internal_Review_Report_Template_v2_0.xlsx")

NAVY, TEAL, GREY, BYLINE = "17365D", "0F6B78", "666666", "808080"
BYLINE_TEXT = "Created by Rowena Chen and Thai Phan"
F = "Calibri"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# status -> (fill, font colour). Same palette everywhere a status appears.
PALETTE = {
    "Exception": ("F4CCCC", "C00000"), "NOT READY": ("F4CCCC", "C00000"),
    "Does Not Tie": ("F4CCCC", "C00000"), "Missing Information": ("F4CCCC", "C00000"),
    "Warning": ("FFF2CC", "7F6000"), "READY WITH COMMENTS": ("FFF2CC", "7F6000"),
    "Reviewed / Monitor": ("DDEBF7", "1F4E78"), "Monitor": ("DDEBF7", "1F4E78"),
    "Supported Elsewhere": ("DDEBF7", "1F4E78"),
    "Pass": ("E2EFDA", "375623"), "Clean": ("E2EFDA", "375623"), "Ties": ("E2EFDA", "375623"),
    "READY FOR RELEASE": ("E2EFDA", "375623"),
}
# The manager's readiness decision. Only a person sets this; the engine leaves it blank.
DECISIONS = ["Ready for Release", "Ready With Comments", "Not Ready"]

# The 16 findings columns, in display order. Must match rowOf() in script 04.
COLUMNS = [
    ("Status", 20), ("Account / Item", 38), ("Finding", 64), ("Required Action", 46),
    ("Difference / Variance / Age", 18), ("Source Value", 16), ("Comparison Value", 16),
    ("Source 1", 30), ("Source 2", 24), ("Test Performed", 52), ("Category", 20), ("Finding Type", 14),
    ("Rule ID", 12), ("Result ID", 28), ("Property Code", 14), ("Review Period", 14),
]
HIDDEN = {"Rule ID", "Result ID", "Property Code", "Review Period"}
WRAP = {"Account / Item", "Finding", "Required Action", "Test Performed", "Source 1", "Source 2"}

# (tab, table, title, description). Order here is the tab order after Executive Summary.
AREA_TABS = [
    ("BSSS", "tblBSSS", "BALANCE SHEET SUPPORT (BSSS)",
     "Each BSSS schedule tied to the Trial Balance, balance sheet accounts with no support, and AR Other / AP Other open items over 30 days."),
    ("Bank Rec", "tblBankRec", "BANK RECONCILIATIONS",
     "Bank rec tie-out to the Trial Balance and the age of every outstanding item. Monitor over 30 days, Warning over 60, Exception over 90."),
    ("T12", "tblT12", "T12 TRENDS",
     "Current month against prior months: utilities, recurring contracts, depreciation and all other income and expense accounts."),
    ("Budget Comparison", "tblBudget", "BUDGET COMPARISON",
     "Major expenses recorded as budgeted, month and year-to-date variances, and agreement to the Trial Balance."),
    ("General Ledger", "tblGL", "GENERAL LEDGER",
     "GL agreement to the Trial Balance, accrual clear-down, and manual journal entries for review."),
    ("PM Fee", "tblPMFee", "PROPERTY MANAGEMENT FEE",
     "Fee base, recalculation, GL tie, hand-entered adjustments traced to journal entries, and revenue spikes."),
    ("Other Controls", "tblOther", "BALANCE SHEET & OTHER CONTROLS",
     "Trial Balance and Balance Sheet integrity, mortgage and escrow, subledger (AME) tie-outs, security deposits, clearing accounts and abnormal balances."),
]
AREA_ROWS = [  # Executive Summary Review Areas table: (area name as the engine writes it, tab to open)
    ("Source Files", "Source Validation"), ("BSSS", "BSSS"), ("Bank Reconciliations", "Bank Rec"),
    ("T12", "T12"), ("Budget Comparison", "Budget Comparison"), ("General Ledger", "General Ledger"),
    ("PM Fee", "PM Fee"), ("Balance Sheet & Other", "Other Controls"),
]
AREA_CAPACITY, ALL_CAPACITY = 100, 500


def font(**kw):
    base = dict(name=F, size=11)
    base.update(kw)
    return Font(**base)


def band(ws, rng, text, size=11):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = text
    c.font = font(size=size, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", indent=1)


def status_formatting(ws, rng):
    for text, (fill, fg) in PALETTE.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                                      fill=PatternFill("solid", fgColor=fill), font=Font(bold=True, color=fg)))


def decision_list(rng):
    """Dropdown for the manager's readiness decision (Executive Summary D8, Review History column D)."""
    dv = DataValidation(type="list", formula1="ManagerDecision", allow_blank=True)  # named range on the hidden Setup sheet
    dv.promptTitle, dv.prompt = "Manager decision", "Choose after reviewing the findings."
    dv.showInputMessage = True
    dv.add(rng)
    return dv


def header_row(ws, row, labels, first_col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row, first_col + i, text)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)


def findings_tab(ws, table_name, title, description, header_at, capacity, lead=()):
    """lead = extra leading columns, e.g. All Results' FindingKey and Review Area."""
    cols = list(lead) + COLUMNS
    last = openpyxl.utils.get_column_letter(len(cols))
    ws.merge_cells(f"A1:{openpyxl.utils.get_column_letter(min(len(cols), 12 + len(lead)))}1")
    ws["A1"].value = title
    ws["A1"].font = font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws["A2"].value = description
    ws["A2"].font = font(size=9, italic=True, color=GREY)
    visible_last = openpyxl.utils.get_column_letter(12 + len(lead))
    ws[f"{visible_last}2"].value = BYLINE_TEXT
    ws[f"{visible_last}2"].font = font(size=8, color=BYLINE)
    ws[f"{visible_last}2"].alignment = Alignment(horizontal="right")
    header_row(ws, header_at, [c[0] for c in cols])
    ws.row_dimensions[header_at].height = 30
    for i, (name, width) in enumerate(cols, start=1):
        L = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[L].width = width
        if name in HIDDEN or name == "FindingKey":
            ws.column_dimensions[L].hidden = True
        for r in range(header_at + 1, header_at + 1 + capacity):
            ws.cell(r, i).alignment = Alignment(vertical="top", wrap_text=name in WRAP)
            ws.cell(r, i).font = font()
    t = Table(displayName=table_name, ref=f"A{header_at}:{last}{header_at + capacity}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    status_col = openpyxl.utils.get_column_letter(len(lead) + 1)
    status_formatting(ws, f"{status_col}{header_at + 1}:{status_col}{header_at + capacity + 2000}")
    ws.freeze_panes = f"{openpyxl.utils.get_column_letter(len(lead) + 3)}{header_at + 1}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


def build_executive_summary(ws):
    for L, w in zip("ABCDEFGH", [34, 16, 15, 15, 15, 15, 15, 44]):
        ws.column_dimensions[L].width = w
    band(ws, "A1:H1", "ACCOUNTING INTERNAL REVIEW REPORT", size=20)
    ws.row_dimensions[1].height = 34
    ws["H2"].value = BYLINE_TEXT
    ws["H2"].font = font(size=8, color=BYLINE)
    ws["H2"].alignment = Alignment(horizontal="right")

    # Header block. Engine writes B3, B4, B5, F3, F4, F5, F6.
    labels = [(3, "Property Code", "Review Attempt"), (4, "Property Name", "Submission Type"),
              (5, "Review Period", "Review Date"), (6, "Required Sources Found", "Report File")]
    for r, left, right in labels:
        ws.cell(r, 1, left).font = font(bold=True, color=GREY)
        ws.cell(r, 5, right).font = font(bold=True, color=GREY)
        ws.merge_cells(f"B{r}:D{r}")
        ws.merge_cells(f"F{r}:H{r}")
        ws.cell(r, 2).font = font()
        ws.cell(r, 6).font = font()
    ws["B6"].value = ("=COUNTIFS('Source Validation'!$B$4:$B$10,\"Required\",'Source Validation'!$F$4:$F$10,\"Found\")"
                      "&\" of \"&COUNTIF('Source Validation'!$B$4:$B$10,\"Required\")")

    # Manager decision. The engine never judges readiness: D8 is left blank for the manager to pick.
    # The four counts below are formulas over All Results (Status is column C).
    band(ws, "A8:C8", "MANAGER DECISION")
    ws.merge_cells("D8:H8")
    ws["D8"].font = font(size=14, bold=True)
    ws["D8"].alignment = Alignment(horizontal="center", vertical="center")
    ws["D8"].border = BOX
    ws.row_dimensions[8].height = 26
    status_formatting(ws, "D8")
    ws.add_data_validation(decision_list("D8"))
    ws["A9"].value = "Chosen by the reviewing manager after reading the findings. The review does not decide readiness."
    ws["A9"].font = font(size=9, italic=True, color=GREY)
    counts = [("A", "B", "Exceptions", "Exception", "F4CCCC", "C00000"),
              ("C", "D", "Warnings", "Warning", "FFF2CC", "7F6000"),
              ("E", "F", "Reviewed / Monitor", "Reviewed / Monitor", "DDEBF7", "1F4E78"),
              ("G", "H", "Passed", "Pass", "E2EFDA", "375623")]
    for a, b, label, status, fill, fg in counts:
        ws.merge_cells(f"{a}10:{b}10")
        ws.merge_cells(f"{a}11:{b}12")
        ws[f"{a}10"].value = label
        ws[f"{a}11"].value = f"=COUNTIF('All Results'!$C:$C,\"{status}\")"
        for cell, size in ((f"{a}10", 11), (f"{a}11", 20)):
            ws[cell].font = font(size=size, bold=True, color=fg)
            ws[cell].fill = PatternFill("solid", fgColor=fill)
            ws[cell].alignment = Alignment(horizontal="center", vertical="center")

    # Review Areas (formulas over All Results: Review Area is column B, Status column C).
    band(ws, "A14:H14", "REVIEW AREAS")
    header_row(ws, 15, ["Area", "Status", "Exceptions", "Warnings", "Monitor", "Passed", "Open"])
    ws.merge_cells("G15:H15")
    for i, (area, tab) in enumerate(AREA_ROWS):
        r = 16 + i
        ws.cell(r, 1, area).font = font(bold=True)
        for col, status in zip("CDEF", ["Exception", "Warning", "Reviewed / Monitor", "Pass"]):
            ws[f"{col}{r}"] = f"=COUNTIFS('All Results'!$B:$B,$A{r},'All Results'!$C:$C,\"{status}\")"
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
        ws[f"B{r}"] = (f'=IF(C{r}>0,"Exception",IF(D{r}>0,"Warning",IF(E{r}>0,"Monitor",'
                       f'IF(F{r}>0,"Clean","Not run"))))')
        ws[f"B{r}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"G{r}:H{r}")
        ws[f"G{r}"] = f"=HYPERLINK(\"#'{tab}'!A1\",\"{tab} tab\")"
        ws[f"G{r}"].font = font(color="0563C1", underline="single")
        for c in "ABCDEFGH":
            ws[f"{c}{r}"].border = BOX
    status_formatting(ws, "B16:B23")

    # BSSS section. Counts are formulas over the BSSS tab (Status = column A, Rule ID = hidden column M).
    band(ws, "A25:H25", "BALANCE SHEET SUPPORT (BSSS)")
    boxes = [("A", "B", "Ties", '=COUNTIFS(BSSS!$M:$M,"BSSS-01",BSSS!$A:$A,"Pass")', "E2EFDA", "375623"),
             ("C", "D", "Does Not Tie", '=COUNTIFS(BSSS!$M:$M,"BSSS-01",BSSS!$A:$A,"Exception")+COUNTIFS(BSSS!$M:$M,"BSSS-03",BSSS!$A:$A,"Exception")', "F4CCCC", "C00000"),
             ("E", "F", "Missing Information", '=COUNTIFS(BSSS!$M:$M,"BSSS-02",BSSS!$A:$A,"Exception")+COUNTIFS(BSSS!$M:$M,"BSSS-02",BSSS!$A:$A,"Warning")', "F4CCCC", "C00000"),
             ("G", "H", "Supported Elsewhere", '=COUNTIFS(BSSS!$M:$M,"BSSS-02",BSSS!$A:$A,"Reviewed / Monitor")', "DDEBF7", "1F4E78")]
    for a, b, label, formula, fill, fg in boxes:
        ws.merge_cells(f"{a}26:{b}26")
        ws.merge_cells(f"{a}27:{b}27")
        ws[f"{a}26"].value, ws[f"{a}27"].value = label, formula
        for cell, size in ((f"{a}26", 10), (f"{a}27", 16)):
            ws[cell].font = font(size=size, bold=True, color=fg)
            ws[cell].fill = PatternFill("solid", fgColor=fill)
            ws[cell].alignment = Alignment(horizontal="center", vertical="center")
    ws["A28"].value = "Supported Elsewhere means the balance is supported by another report (bank rec, AR aging, mortgage statement) rather than a BSSS schedule."
    ws["A28"].font = font(size=9, italic=True, color=GREY)
    header_row(ws, 29, ["Accounts Needing Attention", "Status", "BSSS", "Trial Balance", "Difference", "Finding"])
    ws.merge_cells("F29:H29")
    for r in range(30, 40):  # engine writes A:E as a block and F alone (F:H is merged)
        ws.merge_cells(f"F{r}:H{r}")
        for c in "ABCDEF":
            ws[f"{c}{r}"].font = font()
            ws[f"{c}{r}"].alignment = Alignment(vertical="top", wrap_text=c == "F")
        for c in "CDE":
            ws[f"{c}{r}"].number_format = "#,##0.00;(#,##0.00)"
    status_formatting(ws, "B30:B39")

    # Bank reconciliations. Engine writes A44:H51, one row per reconciliation.
    band(ws, "A41:H41", "BANK RECONCILIATIONS")
    ws["A42"].value = "Clean: nothing outstanding over 30 days  ·  Monitor: 31–60 days  ·  Warning: over 60 days  ·  Exception: over 90 days, or the rec does not tie to the Trial Balance"
    ws["A42"].font = font(size=9, italic=True, color=GREY)
    header_row(ws, 43, ["Bank Account", "Status", "GL Account", "Reconciled Bank", "GL Balance",
                        "Unreconciled Difference", "Oldest Item (days)", "Items Over 30 Days"])
    ws.row_dimensions[43].height = 30
    for r in range(44, 52):
        for c in "ABCDEFGH":
            ws[f"{c}{r}"].font = font()
        for c in "DEF":
            ws[f"{c}{r}"].number_format = "#,##0.00;(#,##0.00)"
        for c in "GH":
            ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
    status_formatting(ws, "B44:B51")

    band(ws, "A53:H53", "MANAGER CONCLUSION")
    ws["A54"].value = "Conclusion / Key Matters:"
    ws["A54"].font = font(bold=True, color=GREY)
    ws.merge_cells("A55:H60")
    ws["A55"].alignment = Alignment(vertical="top", wrap_text=True)
    ws["A55"].border = BOX
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


def extend_source_validation(ws):
    """Keep the v1.2 grid (rows 3-10, engine writes C:F and H of rows 4-10) and add the source findings table."""
    for L, w in zip("ABCDEFGH", [30, 30, 64, 46, 18, 16, 30, 30]):
        ws.column_dimensions[L].width = w
    band(ws, "A12:L12", "SOURCE FINDINGS")
    findings_header = 13
    header_row(ws, findings_header, [c[0] for c in COLUMNS])
    for i, (name, width) in enumerate(COLUMNS, start=1):
        L = openpyxl.utils.get_column_letter(i)
        if i > 8:
            ws.column_dimensions[L].width = width
        if name in HIDDEN:
            ws.column_dimensions[L].hidden = True
        for r in range(findings_header + 1, findings_header + 1 + 30):
            ws.cell(r, i).alignment = Alignment(vertical="top", wrap_text=name in WRAP)
    t = Table(displayName="tblSourceFindings", ref=f"A{findings_header}:P{findings_header + 30}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    status_formatting(ws, f"A{findings_header + 1}:A{findings_header + 200}")
    status_formatting(ws, "F4:F10")


def main():
    wb = openpyxl.load_workbook(SRC)
    for name in ["Exceptions", "Warnings", "Reviewed_Monitor", "Passed Controls", "Executive Summary"]:
        del wb[name]
    es = wb.create_sheet("Executive Summary", 0)
    build_executive_summary(es)
    for tab, table, title, desc in AREA_TABS:
        findings_tab(wb.create_sheet(tab), table, title, desc, header_at=3, capacity=AREA_CAPACITY)
    extend_source_validation(wb["Source Validation"])
    rh = wb["Review History"]  # engine writes one row per attempt; column D is the manager's decision
    rh["D3"].value = "Manager Decision"
    rh.tables["tblReviewHistory"].tableColumns[3].name = "Manager Decision"  # must match the header cell
    setup = wb["Setup"]
    for i, v in enumerate(["ManagerDecision"] + DECISIONS + [None]):
        setup.cell(4, 1 + i).value = v
    del wb.defined_names["OverallStatus"]
    wb.defined_names["ManagerDecision"] = DefinedName("ManagerDecision", attr_text="Setup!$B$4:$D$4")
    rh.data_validations.dataValidation = [d for d in rh.data_validations.dataValidation if str(d.sqref) != "D4:D53"]  # v1.2's OverallStatus list
    rh.add_data_validation(decision_list("D4:D53"))
    status_formatting(rh, "D4:D53")
    allr = wb.create_sheet("All Results")
    findings_tab(allr, "tblAllResults", "ALL RESULTS",
                 "Every test in one table. Filter by Review Area or Status. FindingKey (hidden, column A) is what re-reviews compare on.",
                 header_at=3, capacity=ALL_CAPACITY, lead=(("FindingKey", 30), ("Review Area", 22)))
    order = ["Executive Summary"] + [t[0] for t in AREA_TABS] + ["Source Validation", "Changes Since Last Review",
             "Review History", "All Results", "Setup", "Rule Library"]
    wb._sheets = [wb[n] for n in order]
    for n in ("Setup", "Rule Library"):
        wb[n].sheet_state = "hidden"
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Executive Summary"
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
